"""
无锡太湖学院法学本科毕业论文审核程序
文件名：thesis_law_reviewer.py
功能：使用DeepSeek大模型按照法学论文要求评审论文（支持三种类型）
"""

import os
import json
import re
import time
from docx import Document
from datetime import datetime
import requests
from typing import Dict, List, Tuple, Optional
import logging



# ---- 新口径同步 (2026-09-21): 四形态对象认定 + 审计条款A-F + 确定性预检 ----
import os as _os, sys as _sys
_sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", "增强0908"))
try:
    from audit_principles import PRINCIPLES_TEXT, precheck as _precheck, evidence_block as _evidence_block
except Exception:
    PRINCIPLES_TEXT, _precheck, _evidence_block = "", None, None

OBJ_RULE = ("【研究对象认定(2026-09-20口径)】研究对象四形态均有效: ①具体名称前置 ②脱敏表述前置"
            "('某高校''某企业''某公司''A公司''B集团'等) ③具体名称副标题'——以XX为例' ④脱敏副标题"
            "'——以某XX为例'。只判定是否存在可锚定的研究对象, 不得因脱敏/不具名/位置形态扣分或判不合格。")

def _pa_block(file_path):
    """审计条款+实测证据prompt块(模块未加载时返回空串, 不影响原流程)。"""
    if _precheck is None or not file_path:
        return ""
    try:
        return PRINCIPLES_TEXT + "\n## 实测证据(程序预检, 优先于目测)\n" + _evidence_block(_precheck(file_path))
    except Exception:
        return ""


def _count_refs(thesis_info):
    """程序实测参考文献条数: 优先数'参考文献'章节内 [数字] 编号行, fallback全文统计。"""
    sections = thesis_info.get("sections", {})
    for k, v in sections.items():
        if "参考文献" in k:
            n = sum(1 for ln in v.split("\n") if re.search(r"\[\d+\]", ln))
            if n:
                return n
    return thesis_info.get("stats", {}).get("reference_count", 0)


class ThesisLawReviewer:
    """
    使用DeepSeek的法学论文审核程序
    支持三种类型：案例分析型、规范分析型、案例评释型
    """
    
    def __init__(self, api_key: str = None):
        """
        初始化审核器
        :param api_key: DeepSeek API密钥
        """
        # 设置日志
        self.logger = self._setup_logger()
        
        # DeepSeek API配置
        self.api_key = api_key or os.environ.get("DEEPSEEK_API_KEY", "")  # 2026-09-21修: 读env(百炼), 原硬编码官方key+百炼端点=401
        self.api_base = os.environ.get("DEEPSEEK_BASE_URL", "https://api.deepseek.com/v1")
        self.model = os.environ.get("DEEPSEEK_MODEL", "deepseek-chat")
        
        if not self.api_key:
            self.logger.warning("未设置API密钥，将使用规则进行基础审核")
        
        # 法学论文审核要求
        self.requirements = {
            "通用要求": {
                "正文字数": 8000,
                "摘要字数": (300, 500),
                "关键词数量": (3, 5),
                "参考文献数量": 10
            },
            "三种类型": {
                "案例分析型": {
                    "核心特征": "以一个具体案例为分析对象，围绕争议焦点进行法律分析",
                    "核心章节": ["引言", "案件事实", "争议焦点", "分析", "结论"],
                    "题目要求": "必须体现具体案例名称",
                    "案例数量": 1,
                    "案例位置": "独立章节"
                },
                "规范分析型": {
                    "核心特征": "通过3-5个典型案例揭示普遍性问题，进行理论分析和制度完善",
                    "核心章节": ["引言", "概念与理论", "问题分析", "建议", "结语"],
                    "题目要求": "体现研究问题",
                    "案例数量": "3-5个",
                    "案例位置": "引言中"
                },
                "案例评释型": {
                    "核心特征": "采用请求权基础分析法对具体案例进行鉴定式分析",
                    "核心章节": ["引言", "案件事实", "请求权基础分析", "请求权检视", "结论"],
                    "题目要求": "体现具体案例名称",
                    "案例数量": 1,
                    "案例位置": "独立章节"
                }
            }
        }
        
        # 必达项
        self.must_pass_items = [
            "题目合格",
            "字数合格",
            "框架合格",
            "语言合格",
            "案例来源合格",
            "参考文献合格"
        ]
    
    def _setup_logger(self) -> logging.Logger:
        """设置日志记录器"""
        logger = logging.getLogger("ThesisLawReviewer")
        logger.setLevel(logging.INFO)
        
        # 控制台处理器
        ch = logging.StreamHandler()
        ch.setLevel(logging.INFO)
        
        # 格式器
        formatter = logging.Formatter('%(asctime)s - %(levelname)s - %(message)s')
        ch.setFormatter(formatter)
        
        logger.addHandler(ch)
        
        return logger
    
    def call_deepseek_api(self, messages: List[Dict], temperature: float = 0.3, max_tokens: int = 5000) -> Dict:
        """
        调用DeepSeek API
        """
        headers = {
            "Authorization": f"Bearer {self.api_key}",
            "Content-Type": "application/json"
        }
        
        data = {
            "model": self.model,
            "enable_thinking": False,
            "messages": messages,
            "temperature": temperature,
            "max_tokens": max_tokens,
            "response_format": {"type": "json_object"}
        }
        
        import time as _time
        for _attempt in range(3):
            try:
                response = requests.post(
                    f"{self.api_base}/chat/completions",
                    headers=headers,
                    json=data,
                    timeout=90
                )
                if response.status_code == 200:
                    return response.json()
                self.logger.error(f"API调用失败(第{_attempt+1}次)：{response.status_code} - {response.text[:120]}")
            except Exception as e:
                self.logger.error(f"API调用异常(第{_attempt+1}次)：{str(e)[:120]}")
            if _attempt < 2:
                _time.sleep(2 * (_attempt + 1))
        return None
    
    def extract_text_from_docx(self, file_path: str) -> Tuple[str, Dict]:
        """
        从docx文件中提取文本和结构信息
        """
        try:
            doc = Document(file_path)
            full_text = []
            sections = {}
            current_section = "开头"
            section_content = []
            
            # 统计信息
            stats = {
                "paragraph_count": 0,
                "table_count": len(doc.tables),
                "word_count": 0,
                "reference_count": 0
            }
            
            # 提取所有段落
            for para in doc.paragraphs:
                text = para.text.strip()
                
                if text:
                    full_text.append(text)
                    stats["paragraph_count"] += 1
                    
                    # 检测是否是章节标题
                    if self._is_section_title(text):
                        if section_content:
                            sections[current_section] = "\n".join(section_content)
                        current_section = text
                        section_content = []
                    else:
                        section_content.append(text)
                    
                    # 粗略统计参考文献（包含方括号数字的行）
                    if re.search(r'\[\d+\]', text):
                        stats["reference_count"] += 1
            
            # 添加最后一个章节
            if section_content:
                sections[current_section] = "\n".join(section_content)
            
            # 计算字数
            full_text_str = "\n".join(full_text)
            stats["word_count"] = len(full_text_str)
            
            return full_text_str, {
                "sections": sections,
                "stats": stats,
                "file_name": os.path.basename(file_path),
                "full_text": full_text_str
            }
            
        except Exception as e:
            self.logger.error(f"读取文档失败：{str(e)}")
            return "", {}
    
    def _is_section_title(self, text: str) -> bool:
        """判断是否为章节标题"""
        # 2026-09-21修: ①目录行(标题后跟\t页码/多空格页码)不当标题 ②支持"第1章"阿拉伯数字
        if re.search(r'\t\s*\d{1,3}$', text) or re.search(r'\s{2,}\d{1,3}$', text):
            return False
        patterns = [
            r'^第[一二三四五六七八九十\d]+章',
            r'^摘要$|^abstract$|^目录$|^绪论$|^引言$|^结论$|^参考文献$|^致谢$',
            r'^\d+\.\d+\s+',
            r'^一、|^二、|^三、|^四、|^五、|^六、',
            r'^（一）|^（二）|^（三）'
        ]
        
        for pattern in patterns:
            if re.search(pattern, text, re.IGNORECASE):
                return True
        return False
    
    def extract_title_from_doc(self, doc) -> str:
        """从文档中提取题目(2026-09-21修: 优先解析'论文题目：xxx'标签行, 修复误抓封面校名行)"""
        try:
            # ① 显式标签行: "论文题目：xxx" / "题 目：xxx"(含副标题全长)
            for para in doc.paragraphs[:40]:
                t = para.text.strip()
                m = re.match(r"^论文题目[:：]\s*(.+)$", t) or re.match(r"^题\s*目[:：]\s*(.+)$", t)
                if m and m.group(1).strip():
                    return m.group(1).strip()
            # ② 启发式fallback: 排除校名/学位/答辩类封面行
            skip = re.compile(r"学院|大学|自学考试|本科|专科|专升本|答辩|专业|姓名|学号|指导|日期|年|月")
            for para in doc.paragraphs[:30]:
                t = para.text.strip()
                if 12 <= len(t) <= 60 and not skip.search(t) and "摘" not in t:
                    return t
        except:
            pass
        return "未找到题目"
    
    def check_word_count(self, word_count: int) -> Tuple[bool, str]:
        """检查字数"""
        if word_count >= 8000:
            return True, f"字数达标：{word_count}字"
        else:
            return False, f"字数不足：{word_count}字，要求不少于8000字"
    
    def detect_thesis_type(self, sections: Dict, title: str) -> Tuple[str, List[str]]:
        """
        根据章节结构和题目判断论文类型
        返回：(类型, 存在的核心章节)
        """
        section_names = " ".join(list(sections.keys())).lower()
        
        # 检查题目是否包含案例名称
        has_case_in_title = bool(re.search(r'案|例|纠纷|诉', title))
        
        # 检查核心章节
        found_sections = []
        
        # 案例分析型核心章节关键词
        case_analysis_keywords = ["案件事实", "基本事实", "案情", "争议焦点", "焦点"]
        for keyword in case_analysis_keywords:
            if keyword in section_names:
                found_sections.append(keyword)
        
        # 规范分析型核心章节关键词
        normative_keywords = ["概念", "理论", "问题分析", "存在问题", "完善建议", "对策"]
        for keyword in normative_keywords:
            if keyword in section_names:
                found_sections.append(keyword)
        
        # 案例评释型核心章节关键词
        review_keywords = ["请求权", "鉴定", "基础分析", "检视", "要件"]
        for keyword in review_keywords:
            if keyword in section_names:
                found_sections.append(keyword)
        
        # 判断类型
        if "争议焦点" in section_names or "案件事实" in section_names:
            if has_case_in_title:
                return "案例分析型", found_sections
            else:
                return "案例分析型", found_sections
        elif "请求权" in section_names or "检视" in section_names:
            return "案例评释型", found_sections
        else:
            return "规范分析型", found_sections
    
    def review_with_deepseek(self, thesis_text: str, thesis_info: Dict, title: str, thesis_type: str) -> Dict:
        """
        使用DeepSeek审核法学论文
        """
        stats = thesis_info.get("stats", {})
        sections = thesis_info.get("sections", {})
        word_count = stats.get("word_count", 0)
        
        # 获取该类型的核心章节要求
        type_requirements = self.requirements["三种类型"].get(thesis_type, {})
        required_chapters = type_requirements.get("核心章节", [])
        
        # 检查现有章节
        section_names = list(sections.keys())
        
        # 构建Prompt
        _pa = _pa_block(thesis_info.get('_source_path'))
        _refs_n = _count_refs(thesis_info)
        prompt = f"""
        你是一位严格的本科法学毕业论文评审专家。请按照以下要求对论文进行详细评审。

        ## 论文基本信息
        - 题目：{title}
        - 论文类型：{thesis_type}
        - 总字数：{word_count}字
        - 章节列表：{section_names}
        - 参考文献条数（程序实测）：{_refs_n} 条 —— 该数字由程序从文档参考文献区实统计, 是判断"参考文献数量"的唯一依据, 你不得臆断数量不足或数量达标, 只依据此数。

        ## 审核要求

        ### 从严原则（最高优先级, 适用于内容与结构判定）
        0. 内容结构从严、逻辑从严：对论文的内容结构、章节逻辑、论证逻辑问题一律从严判定, 宁错勿放——
           ① 存在与论文主题无关的实质章节/大段内容（如正当防卫论文中整节论述"算法从属性"等无关内容）属严重错误, 框架必判不合格;
           ② 核心章节严重缺失/缺失两个及以上核心章节 → 框架不合格;
           ③ 章节编号重复、标题错位、内容错位（某节内容与标题不符）→ 框架/结构从严, 不得以"小瑕疵"放行;
           ④ 论证逻辑断裂、问题-建议不对应等逻辑问题 → 论证分析从严扣分, 逻辑严重混乱者语言/框架项不合格;
           ⑤ 把握不准时从严判定（宁错勿放）：结构或逻辑存在疑似的严重问题时, 判不合格; 分项评分如实反映严重程度。

        ### 必达项（必须全部达到，否则不合格）
        1. 题目合格：{type_requirements.get('题目要求', '符合类型要求')}
           案例名称/制度载体的脱敏表述(如"某公司劳动争议案""XX案")同四形态有效, 不得因脱敏扣分
        2. 字数合格：正文字数不少于8000字（以程序实测总字数为准, 不必人工重数）
        3. 框架合格：基本符合所选类型的框架结构，包含核心章节；同时必须无重大结构缺陷（无关章节、重大缺失、章节编号重复、内容错位）, 见从严原则
        4. 语言合格：语句通顺专业，逻辑合理，观点明确；逻辑问题从严见从严原则
        5. 案例来源合格：案例真实，标注来源（裁判文书网/北大法宝等）
        6. 参考文献合格：程序实测 {_refs_n} 条, 达到10条即数量合格（此条由程序判定, 你只需评质量格式）

        ### {thesis_type}核心章节要求
        应包含以下核心章节（标题可以略有不同）：
        {json.dumps(required_chapters, ensure_ascii=False, indent=2)}

        ## 论文内容（前5000字）
        {thesis_text[:5000]}

        {_pa}

        
## 评审报告撰写规范(最高优先级, 适用于所有评语与"评审报告"字段)
1. 评语以自然语言撰写, 像一位毕业论文指导教师在评阅: 禁止"实测""检测到""程序判定""证据""预检""指标显示"等程序化词汇;
   数据表述用自然形式(如"全文约1.4万字, 插图14幅、表格3张")。
2. JSON新增顶层字段 "评审报告"(字符串): 一篇完整评语, 严格按三部分组织, 每部分以标记行开头:
【总体评价】2-4句: 论文整体质量与结论。若合格, 写明建议成绩(如"综合评定为合格, 建议成绩82分");
  若不合格, 只作定性结论, 不得出现任何分数、分值、得分。
3. 【具体评价】(在"评审报告"内): 按题目、结构、内容与论证、语言与规范等维度逐段展开,
   所有具体问题在此详细指出(引到章节、数值、图表编号等原文位置), 用连贯段落, 不用表格罗列。
4. 【修改意见】(在"评审报告"内): 逐条列出, 具体可执行; 不合格论文的修改意见要覆盖全部主要缺陷。
5. 其余分项字段照常输出(供内部判定), 但"评审报告"是最终呈现文本, 必须自足完整。

请按以下JSON格式输出评审结果：

        {{
            "是否合格": true/false,
  "评审报告": "【总体评价】...【具体评价】...【修改意见】...",
            "总分": 60-100之间的整数,
            "必达项检查": {{
                "题目合格": true/false,
                "字数合格": true/false,
                "框架合格": true/false,
                "语言合格": true/false,
                "案例来源合格": true/false,
                "参考文献合格": true/false
            }},
            "各项评分": {{
                "题目": {{
                    "得分": 0-15,
                    "评语": "详细评价",
                    "问题": []
                }},
                "字数": {{
                    "得分": 0-10,
                    "评语": "详细评价",
                    "问题": []
                }},
                "框架结构": {{
                    "得分": 0-20,
                    "评语": "评价框架是否完整、逻辑是否正确",
                    "包含的核心章节": [],
                    "缺失的核心章节": [],
                    "问题": []
                }},
                "论证分析": {{
                    "得分": 0-20,
                    "评语": "评价论证是否充分、逻辑是否严密",
                    "问题": []
                }},
                "案例运用": {{
                    "得分": 0-15,
                    "评语": "评价案例是否真实、运用是否恰当、来源是否标注",
                    "问题": []
                }},
                "参考文献": {{
                    "得分": 0-10,
                    "评语": "评价参考文献数量、质量、格式",
                    "问题": []
                }},
                "语言表达": {{
                    "得分": 0-10,
                    "评语": "评价语言是否专业通顺",
                    "问题": []
                }}
            }},
            "总体评语": "详细总结论文优缺点",
            "修改建议": ["建议1", "建议2", "建议3", "建议4"]
        }}
        """
        
        messages = [
            {"role": "system", "content": f"你是严格的法学本科毕业论文评审专家。论文类型为{thesis_type}，请按该类型的要求进行评审。"},
            {"role": "user", "content": prompt}
        ]
        
        response = self.call_deepseek_api(messages, max_tokens=5000)
        
        if response and 'choices' in response:
            try:
                result = json.loads(response['choices'][0]['message']['content'])
                return result
            except Exception as e:
                self.logger.error(f"解析API返回结果失败：{str(e)}")
        
        # API调用失败，使用规则审核
        return self._rule_based_review(thesis_text, thesis_info, title, thesis_type)
    
    def _rule_based_review(self, thesis_text: str, thesis_info: Dict, title: str, thesis_type: str) -> Dict:
        """
        基于规则的备用审核方案
        """
        stats = thesis_info.get("stats", {})
        sections = thesis_info.get("sections", {})
        word_count = stats.get("word_count", 0)
        
        # 获取该类型的核心章节要求
        type_requirements = self.requirements["三种类型"].get(thesis_type, {})
        required_chapters = type_requirements.get("核心章节", [])
        
        # 检查核心章节
        section_names_lower = " ".join(list(sections.keys())).lower()
        found_chapters = []
        missing_chapters = []
        
        for chapter in required_chapters:
            if chapter.lower() in section_names_lower:
                found_chapters.append(chapter)
            else:
                # 模糊匹配
                found = False
                for keyword in [chapter, chapter[:2]]:
                    if keyword in section_names_lower:
                        found = True
                        found_chapters.append(chapter)
                        break
                if not found:
                    missing_chapters.append(chapter)
        
        # 检查字数
        word_qualified, word_msg = self.check_word_count(word_count)
        
        # 检查题目（案例分析型必须有具体案例名称）
        title_qualified = True
        title_issue = ""
        if thesis_type == "案例分析型" or thesis_type == "案例评释型":
            if not re.search(r'案|例|纠纷|诉', title):
                title_qualified = False
                title_issue = "题目缺少具体案例名称，应体现研究的具体案件"
        
        # 检查参考文献数量
        ref_count = stats.get("reference_count", 0)
        ref_qualified = ref_count >= 10
        
        # 框架合格
        framework_qualified = len(missing_chapters) <= 1  # 允许缺失1个核心章节
        
        # 计算总分
        total_score = 70
        if title_qualified:
            total_score += 5
        if word_qualified:
            total_score += 5
        if framework_qualified:
            total_score += 10
        if ref_qualified:
            total_score += 5
        total_score = max(60, min(95, total_score))
        
        # 必达项
        must_pass = {
            "题目合格": title_qualified,
            "字数合格": word_qualified,
            "框架合格": framework_qualified,
            "语言合格": True,
            "案例来源合格": True,
            "参考文献合格": ref_qualified
        }
        
        is_qualified = all(must_pass.values())
        
        return {
            "是否合格": is_qualified,
            "总分": total_score,
            "必达项检查": must_pass,
            "各项评分": {
                "题目": {"得分": 15 if title_qualified else 8, "评语": title_issue if title_issue else "题目符合要求", "问题": [title_issue] if title_issue else []},
                "字数": {"得分": 10 if word_qualified else 5, "评语": word_msg, "问题": []},
                "框架结构": {"得分": 18 if framework_qualified else 10, "评语": f"包含核心章节：{found_chapters}，缺失：{missing_chapters}", "包含的核心章节": found_chapters, "缺失的核心章节": missing_chapters, "问题": []},
                "论证分析": {"得分": 15, "评语": "论证基本充分", "问题": []},
                "案例运用": {"得分": 12, "评语": "案例运用基本合理", "问题": []},
                "参考文献": {"得分": 8 if ref_qualified else 4, "评语": f"参考文献{ref_count}条", "问题": []},
                "语言表达": {"得分": 8, "评语": "语言基本通顺", "问题": []}
            },
            "总体评语": f"论文{'达到' if is_qualified else '未达到'}基本要求。",
            "修改建议": self._generate_suggestions(must_pass, missing_chapters, thesis_type)
        }
    
    def _generate_suggestions(self, must_pass: Dict, missing_chapters: List[str], thesis_type: str) -> List[str]:
        """生成修改建议"""
        suggestions = []
        
        if not must_pass.get("题目合格"):
            if thesis_type in ["案例分析型", "案例评释型"]:
                suggestions.append("题目必须体现具体案例名称，例如：'XXX案分析'、'XXX案法律问题研究'")
            else:
                suggestions.append("题目应明确研究问题，体现论文核心内容")
        
        if not must_pass.get("字数合格"):
            suggestions.append("正文字数不足8000字，请扩充内容")
        
        if not must_pass.get("框架合格"):
            suggestions.append(f"框架不完整，缺少以下核心章节：{', '.join(missing_chapters)}")
            suggestions.append("请按所选类型的框架要求调整章节结构")
        
        if not must_pass.get("参考文献合格"):
            suggestions.append("参考文献数量不足10条，请补充相关文献")
        
        if not suggestions:
            suggestions.append("论文基本符合要求，可进一步优化论证深度")
        
        return suggestions
    
    def generate_report(self, review_result: Dict, thesis_info: Dict, title: str, thesis_type: str) -> str:
        """
        生成详细的审核报告
        """
        report = []
        report.append("="*70)
        report.append("无锡太湖学院法学本科毕业论文审核报告")
        report.append("="*70)
        report.append("")
        
        # 基本信息
        report.append("一、论文基本信息")
        report.append("-"*50)
        report.append(f"论文题目：{title}")
        report.append(f"论文类型：{thesis_type}")
        report.append(f"论文文件：{thesis_info.get('file_name', '未知')}")
        report.append(f"总字数：{thesis_info.get('stats', {}).get('word_count', 0)}字")
        report.append("")
        
        # 审核结论
        report.append("二、审核结论")
        report.append("-"*50)
        report.append(f"审核结果：{'合 格' if review_result.get('是否合格') else '不合格'}")
        report.append(f"最终得分：{review_result.get('总分', 0)}分")
        report.append("")
        
        # 必达项检查
        report.append("三、必达项检查（全部通过方为合格）")
        report.append("-"*50)
        must_pass = review_result.get('必达项检查', {})
        items = [
            ("题目合格", "1. 题目符合类型要求（案例分析型需体现案例名称）"),
            ("字数合格", "2. 正文字数不少于8000字"),
            ("框架合格", "3. 框架符合所选类型，包含核心章节"),
            ("语言合格", "4. 语言专业通顺，逻辑合理"),
            ("案例来源合格", "5. 案例真实，标注来源"),
            ("参考文献合格", "6. 参考文献不少于10条")
        ]
        for key, desc in items:
            result = must_pass.get(key, False)
            report.append(f"   {'✓' if result else '✗'} {desc}")
        report.append("")
        
        # 分项评分
        report.append("四、分项评分与详细评语")
        report.append("-"*50)
        scores = review_result.get('各项评分', {})
        
        for item, score_info in scores.items():
            report.append(f"\n【{item}】{score_info.get('得分', 0)}分")
            if score_info.get('评语'):
                report.append(f"评语：{score_info['评语']}")
            if item == "框架结构" and score_info.get('缺失的核心章节'):
                report.append(f"缺失的核心章节：{', '.join(score_info['缺失的核心章节'])}")
            if score_info.get('问题'):
                for issue in score_info['问题']:
                    report.append(f"问题：{issue}")
        
        report.append("")
        
        # 总体评语
        report.append("五、总体评语")
        report.append("-"*50)
        report.append(review_result.get('总体评语', '无'))
        report.append("")
        
        # 修改建议
        report.append("六、修改建议")
        report.append("-"*50)
        suggestions = review_result.get('修改建议', [])
        for i, suggestion in enumerate(suggestions, 1):
            report.append(f"{i}. {suggestion}")
        
        report.append("")
        report.append("="*70)
        report.append(f"审核时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        report.append("="*70)
        
        return "\n".join(report)
    
    def save_report(self, report: str, output_path: str = None):
        """保存审核报告"""
        if not output_path:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            output_path = f"法学论文审核报告_{timestamp}.txt"
        
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(report)
        
        self.logger.info(f"审核报告已保存至：{output_path}")
        return output_path
    
    def review_thesis(self, file_path: str, thesis_type: str = None) -> Dict:
        """
        审核单篇法学论文
        :param file_path: 论文文件路径
        :param thesis_type: 论文类型（可选，不指定则自动判断）
        """
        self.logger.info(f"开始审核法学论文：{file_path}")
        
        # 1. 提取论文内容
        thesis_text, thesis_info = self.extract_text_from_docx(file_path)
        if not thesis_text:
            self.logger.error("论文内容提取失败")
            return None
        
        # 2. 提取题目
        doc = Document(file_path)
        title = self.extract_title_from_doc(doc)
        
        # 3. 判断论文类型
        sections = thesis_info.get("sections", {})
        if not thesis_type:
            thesis_type, found = self.detect_thesis_type(sections, title)
            self.logger.info(f"自动判断论文类型为：{thesis_type}")
        else:
            self.logger.info(f"用户指定论文类型为：{thesis_type}")
        
        self.logger.info(f"论文题目：{title}")
        self.logger.info(f"字数：{thesis_info['stats']['word_count']}字")
        
        # 4. 使用DeepSeek审核
        self.logger.info("正在使用DeepSeek进行审核...")
        thesis_info['_source_path'] = file_path
        review_result = self.review_with_deepseek(thesis_text, thesis_info, title, thesis_type)

        # 4b. 程序侧兜底(2026-09-22从严口径):
        # ① 参考文献数量合格由程序硬判(实测条数>=10), 覆盖LLM误判;
        # ② 任一必达项不合格 → 强制 是否合格=false, 防止LLM自相矛盾(如必达项挂但判合格)。
        if review_result:
            _refs_n = _count_refs(thesis_info)
            mp = review_result.get("必达项检查") or {}
            mp["参考文献合格"] = _refs_n >= 10
            review_result["必达项检查"] = mp
            if not all(mp.values()):
                review_result["是否合格"] = False
            self.logger.info(f"程序侧兜底: 实测参考文献{_refs_n}条, 必达项全过={all(mp.values())}, 最终是否合格={review_result['是否合格']}")
        
        # 5. 生成报告
        report = self.generate_report(review_result, thesis_info, title, thesis_type)
        
        # 6. 保存报告
        self.save_report(report)
        
        # 7. 打印结果
        self.logger.info(f"审核完成！")
        self.logger.info(f"结果：{'合格' if review_result.get('是否合格') else '不合格'}")
        self.logger.info(f"得分：{review_result.get('总分')}分")
        
        return {
            "result": review_result,
            "report": report,
            "thesis_info": thesis_info,
            "thesis_type": thesis_type
        }


def main():
    """
    主函数
    """
    print("="*70)
    print("无锡太湖学院法学本科毕业论文审核程序")
    print("（支持案例分析型、规范分析型、案例评释型）")
    print("="*70)
    print()
    print("【审核说明】")
    print("1. 正文字数不少于8000字")
    print("2. 参考文献不少于10条")
    print("3. 案例分析型/案例评释型：题目需体现具体案例名称")
    print("4. 规范分析型：需有3-5个案例在引言中引出问题")
    print("="*70)
    print()
    
    # 使用提供的API密钥
    api_key = os.environ.get("DEEPSEEK_API_KEY", "")
    
    # 获取论文文件路径
    while True:
        file_path = input("请输入论文文件路径（支持.docx格式）：").strip()
        if not file_path:
            print("错误：必须指定论文文件路径")
            continue
        
        if not os.path.exists(file_path):
            print(f"错误：文件不存在 - {file_path}")
            continue
        
        if not file_path.endswith('.docx'):
            print("警告：程序主要支持.docx格式，其他格式可能无法正确读取")
            confirm = input("是否继续？(y/n，默认n)：").strip().lower()
            if confirm != 'y':
                continue
        
        break
    
    # 选择论文类型
    print("\n请选择论文类型（直接回车则自动判断）：")
    print("1. 案例分析型（以一个具体案例为分析对象）")
    print("2. 规范分析型（由案例引出问题，进行制度完善）")
    print("3. 案例评释型（鉴定式案例分析）")
    print("0. 自动判断")
    
    type_choice = input("请输入数字（1/2/3/0，默认0）：").strip()
    thesis_type = None
    if type_choice == "1":
        thesis_type = "案例分析型"
    elif type_choice == "2":
        thesis_type = "规范分析型"
    elif type_choice == "3":
        thesis_type = "案例评释型"
    
    print("\n正在初始化审核程序...")
    
    # 创建审核器
    reviewer = ThesisLawReviewer(api_key=api_key)
    
    # 执行审核
    print("\n开始审核，请稍候（约60秒）...")
    result = reviewer.review_thesis(file_path, thesis_type)
    
    if result:
        print("\n" + "="*70)
        print("审核完成！")
        print(f"论文类型：{result['thesis_type']}")
        print(f"审核结果：{'合 格' if result['result'].get('是否合格') else '不合格'}")
        print(f"最终得分：{result['result'].get('总分')}分")
        print("\n详细报告已保存到文件：法学论文审核报告_时间戳.txt")
        print("="*70)
    else:
        print("\n审核失败，请检查日志。")


if __name__ == "__main__":
    main()