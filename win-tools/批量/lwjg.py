"""
无锡太湖学院本科毕业论文审核程序
文件名：thesis_deepseek_reviewer.py
功能：使用DeepSeek大模型按照学校要求评审论文
"""

import os
import json
import re
import time
from docx import Document
from datetime import datetime
import requests
from typing import Dict, List, Tuple, Optional
import logging



# ---- 新口径同步 (2026-09-21): 四形态对象认定 + 审计条款A-F + 确定性预检 ----
import os as _os, sys as _sys
_sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", "增强0908"))
try:
    from audit_principles import PRINCIPLES_TEXT, precheck as _precheck, evidence_block as _evidence_block
except Exception:
    PRINCIPLES_TEXT, _precheck, _evidence_block = "", None, None

OBJ_RULE = ("【研究对象认定(2026-09-20口径)】研究对象四形态均有效: ①具体名称前置 ②脱敏表述前置"
            "('某高校''某企业''某公司''A公司''B集团'等) ③具体名称副标题'——以XX为例' ④脱敏副标题"
            "'——以某XX为例'。只判定是否存在可锚定的研究对象, 不得因脱敏/不具名/位置形态扣分或判不合格。")

def _pa_block(file_path):
    """审计条款+实测证据prompt块(模块未加载时返回空串, 不影响原流程)。"""
    if _precheck is None or not file_path:
        return ""
    try:
        return PRINCIPLES_TEXT + "\n## 实测证据(程序预检, 优先于目测)\n" + _evidence_block(_precheck(file_path))
    except Exception:
        return ""

class ThesisDeepSeekReviewer:
    """
    使用DeepSeek的论文审核程序
    严格按照无锡太湖学院的要求进行评审
    """
    
    def __init__(self, api_key: str = None):
        """
        初始化审核器
        :param api_key: DeepSeek API密钥
        """
        # 设置日志
        self.logger = self._setup_logger()
        
        # DeepSeek API配置
        self.api_key = api_key or os.environ.get("DEEPSEEK_API_KEY", "")  # 2026-09-21修: 读env(百炼), 原硬编码官方key+百炼端点=401
        self.api_base = os.environ.get("DEEPSEEK_BASE_URL", "https://api.deepseek.com/v1")
        self.model = os.environ.get("DEEPSEEK_MODEL", "deepseek-chat")
        
        if not self.api_key:
            self.logger.warning("未设置API密钥，将使用规则进行基础审核")
        
        # 无锡太湖学院论文要求（严格按照文档）
        self.requirements = {
            "题目要求": "以具体企业、项目、案例为研究对象(脱敏表述同四形态有效, 见OBJ_RULE)",
            "字数要求": 10000,
            "结构要求": {
                "必须包含的部分": [
                    "第一部分：引言（绪论）",
                    "第二部分：相关理论基础（可表述为国内外研究现状等）",
                    "第三部分：研究对象现状",
                    "第四部分：问题和原因分析（可分成两个大章）",
                    "第五部分：解决方案或策略建议",
                    "第六部分：结论"
                ],
                "注意事项": [
                    "每章的标题可以表述上略有不同",
                    "结构逻辑正确的情况下，章节安排可以略有不同",
                    "问题和原因分析可以分成两个大章",
                    "不能发生不同章节事实重复论述同一内容"
                ]
            },
            "语言要求": "语句通顺专业，不能有口语化表达，逻辑合理，能够应用理论进行问题分析并提出解决方案",
            "数据要求": "应有适当的数据或图表做为分析的支撑，图表是对数据分析的可视化呈现",
            "格式要求": "格式基本符合要求，文章本身的上下文格式必须保持一致（封面或标题页格式问题不做要求）"
        }
        
        # 必达项（1/2/3/4/5必须达到，否则不合格）
        self.must_have_items = [
            "题目以具体企业/项目/案例为研究对象",
            "字数达到10000字以上",
            "包含六部分基本结构（标题可略有不同，逻辑正确）",
            "语言专业通顺，逻辑合理，能应用理论分析问题",
            "有适当的数据或图表作为分析支撑"
        ]
        
        # 评分标准
        self.scoring_criteria = {
            "题目": {"weight": 15, "description": "题目是否以具体企业/项目/案例为研究对象"},
            "字数": {"weight": 10, "description": "字数是否达到10000字以上"},
            "结构": {"weight": 25, "description": "结构是否完整，是否符合六部分要求，是否有重复论述"},
            "语言逻辑": {"weight": 20, "description": "语言是否专业通顺，逻辑是否合理，是否能应用理论分析问题"},
            "数据支撑": {"weight": 20, "description": "是否有适当的数据或图表作为分析支撑"},
            "格式": {"weight": 10, "description": "格式是否符合要求，上下文格式一致（封面标题页不做要求）"}
        }
    
    def _setup_logger(self) -> logging.Logger:
        """设置日志记录器"""
        logger = logging.getLogger("ThesisReviewer")
        logger.setLevel(logging.INFO)
        
        # 控制台处理器
        ch = logging.StreamHandler()
        ch.setLevel(logging.INFO)
        
        # 格式器
        formatter = logging.Formatter('%(asctime)s - %(levelname)s - %(message)s')
        ch.setFormatter(formatter)
        
        logger.addHandler(ch)
        
        return logger
    
    def call_deepseek_api(self, messages: List[Dict], temperature: float = 0.3, max_tokens: int = 4000) -> Dict:
        """
        调用DeepSeek API
        """
        headers = {
            "Authorization": f"Bearer {self.api_key}",
            "Content-Type": "application/json"
        }
        
        data = {
            "model": self.model,
            "enable_thinking": False,
            "messages": messages,
            "temperature": temperature,
            "max_tokens": max_tokens,
            "response_format": {"type": "json_object"}
        }
        
        # 2026-09-21修: 间歇401/网络抖动→3次退避重试(端点多实例鉴权状态不同步, 重试有效)
        import time as _time
        for _attempt in range(3):
            try:
                response = requests.post(
                    f"{self.api_base}/chat/completions",
                    headers=headers,
                    json=data,
                    timeout=90
                )
                if response.status_code == 200:
                    return response.json()
                self.logger.error(f"API调用失败(第{_attempt+1}次)：{response.status_code} - {response.text[:120]}")
            except Exception as e:
                self.logger.error(f"API调用异常(第{_attempt+1}次)：{str(e)[:120]}")
            if _attempt < 2:
                _time.sleep(2 * (_attempt + 1))
        return None
    
    def extract_text_from_docx(self, file_path: str) -> Tuple[str, Dict]:
        """
        从docx文件中提取文本和结构信息
        封面或标题页格式问题不做要求，但内容需要提取
        """
        try:
            doc = Document(file_path)
            full_text = []
            sections = {}
            current_section = "开头"
            section_content = []
            
            # 统计信息
            stats = {
                "paragraph_count": 0,
                "table_count": len(doc.tables),
                "image_count": 0,
                "word_count": 0
            }
            
            # 提取所有段落（封面也提取，但格式问题不做要求）
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    full_text.append(text)
                    stats["paragraph_count"] += 1
                    
                    # 检测是否是章节标题
                    if self._is_section_title(text):
                        if section_content:
                            sections[current_section] = "\n".join(section_content)
                        current_section = text
                        section_content = []
                    else:
                        section_content.append(text)
            
            # 添加最后一个章节
            if section_content:
                sections[current_section] = "\n".join(section_content)
            
            # 计算字数
            full_text_str = "\n".join(full_text)
            stats["word_count"] = len(full_text_str)
            
            # 统计图片
            for para in doc.paragraphs:
                # 2026-09-22修: VML老格式(<w:pict>/<v:imagedata>)也认, 不再只认graphicData
                _px = para._element.xml
                stats["image_count"] += _px.count('<a:blip') + _px.count('<v:imagedata')
            # 2026-09-22修: 扫描表格单元格中的图片(与lwsj.py同构修复)——经管论文也常把图
            # 放进表格单元格, 只遍历doc.paragraphs会漏数, 导致"数据支撑不足"误判
            for tb in doc.tables:
                for row in tb.rows:
                    for cell in row.cells:
                        _cx = cell._tc.xml
                        stats["image_count"] += _cx.count('<a:blip') + _cx.count('<v:imagedata')
            
            return full_text_str, {
                "sections": sections,
                "stats": stats,
                "file_name": os.path.basename(file_path)
            }
            
        except Exception as e:
            self.logger.error(f"读取文档失败：{str(e)}")
            return "", {}
    
    def _is_section_title(self, text: str) -> bool:
        """判断是否为章节标题"""
        # 2026-09-21修: ①目录行(标题后跟\t页码/多空格页码)不当标题 ②支持"第1章"阿拉伯数字
        if re.search(r'\t\s*\d{1,3}$', text) or re.search(r'\s{2,}\d{1,3}$', text):
            return False
        patterns = [
            r'^第[一二三四五六七八九十\d]+章',
            r'^摘要$|^abstract$|^目录$|^绪论$|^引言$|^结论$|^参考文献$|^致谢$',
            r'^\d+\.\d+\s+',
            r'^第一部分|^第二部分|^第三部分|^第四部分|^第五部分|^第六部分'
        ]
        
        for pattern in patterns:
            if re.search(pattern, text, re.IGNORECASE):
                return True
        return False
    
    def check_title_format(self, title: str) -> Tuple[bool, str]:
        """
        检查题目格式
        以具体企业、项目、案例为研究对象
        """
        patterns = [
            r'.*公司.*',
            r'.*项目.*',
            r'.*案例.*',
            r'以.*为例',
            r'.*企业.*',
            r'.*厂.*',
            r'.*集团.*'
        ]
        
        for pattern in patterns:
            if re.search(pattern, title):
                return True, "题目包含具体研究对象"
        
        return False, "题目缺少具体研究对象（必须包含企业/项目/案例名称）"
    
    def extract_title_from_doc(self, doc) -> str:
        """从文档中提取题目(2026-09-21修: 优先解析'论文题目：xxx'标签行, 修复误抓封面校名行)"""
        try:
            # ① 显式标签行: "论文题目：xxx" / "题 目：xxx"(含副标题全长)
            for para in doc.paragraphs[:40]:
                t = para.text.strip()
                m = re.match(r"^论文题目[:：]\s*(.+)$", t) or re.match(r"^题\s*目[:：]\s*(.+)$", t)
                if m and m.group(1).strip():
                    return m.group(1).strip()
            # ② 启发式fallback: 排除校名/学位/答辩类封面行
            skip = re.compile(r"学院|大学|自学考试|本科|专科|专升本|答辩|专业|姓名|学号|指导|日期|年|月")
            for para in doc.paragraphs[:30]:
                t = para.text.strip()
                if 12 <= len(t) <= 60 and not skip.search(t) and "摘" not in t:
                    return t
        except:
            pass
        return "未找到题目"
    
    def review_with_deepseek(self, thesis_text: str, thesis_info: Dict, title: str) -> Dict:
        """
        使用DeepSeek审核论文（返回详细评语）
        """
        stats = thesis_info.get("stats", {})
        
        _pa = _pa_block(thesis_info.get('_source_path'))
        prompt = f"""
        你是一位严格的本科毕业论文评审专家。请按照无锡太湖学院的要求，对以下论文进行详细评审，给出详细的评语。

        ## 论文基本信息
        - 题目：{title}
        - 总字数：{stats.get('word_count', 0)}字
        - 表格数量：{stats.get('table_count', 0)}个
        - 图片数量：{stats.get('image_count', 0)}张

        ## 审核要求（1/2/3/4/5必须达到，否则不合格）
        1. 题目要求：以具体企业、项目、案例为研究对象
           例如："财务的成本控制分析以棒杰服装公司为例"、"Z公司销售人员薪酬方案改进研究"
           {OBJ_RULE}
        
        2. 字数要求：10000字以上
        
        3. 结构要求：
           必须包含以下部分（每章的标题可以表述上略有不同）：
           - 第一部分：引言（绪论）
           - 第二部分：相关理论基础（可表述为国内外研究现状等）
           - 第三部分：研究对象现状
           - 第四部分：问题和原因分析（可分成两个大章）
           - 第五部分：解决方案或策略建议
           - 第六部分：结论
           
           注意事项：
           - 结构逻辑必须正确
           - 章节安排可以略有不同
           - 不能发生不同章节事实重复论述同一内容的情况
        
        4. 语言要求：语句通顺专业，不能有口语化表达，逻辑合理，能够应用理论进行问题分析并提出解决方案
        
        5. 数据要求：应有适当的数据或图表做为分析的支撑，图表是对数据分析的可视化呈现
        
        6. 格式要求：格式基本符合要求，文章本身的上下文格式必须保持一致（封面或标题页格式问题不做要求）

        ## 论文内容（前6000字）
        {thesis_text[:6000]}

        {_pa}

        
## 评审报告撰写规范(最高优先级, 适用于所有评语与"评审报告"字段)
1. 评语以自然语言撰写, 像一位毕业论文指导教师在评阅: 禁止"实测""检测到""程序判定""证据""预检""指标显示"等程序化词汇;
   数据表述用自然形式(如"全文约1.4万字, 插图14幅、表格3张")。
2. JSON新增顶层字段 "评审报告"(字符串): 一篇完整评语, 严格按【总体评价】→【具体评价】→【修改意见】三段固定顺序组织,
   三段先后顺序不得颠倒, 且【总体评价】【具体评价】【修改意见】三个标记必须各自另起一行独占段首,
   段与段之间空一行, 任何标记不得接在上一段末尾同一行。
【总体评价】2-4句: 论文整体质量与结论。若合格, 写明建议成绩(如"综合评定为合格, 建议成绩82分");
  若不合格, 只作定性结论, 不得出现任何分数、分值、得分。
3. 【具体评价】(在"评审报告"内): 按题目、结构、内容与论证、语言与规范等维度逐段展开,
   所有具体问题在此详细指出(引到章节、数值、图表编号等原文位置), 用连贯段落, 不用表格罗列。
4. 【修改意见】(在"评审报告"内): 逐条列出, 具体可执行; 不合格论文的修改意见要覆盖全部主要缺陷。
5. 封面页不属于论文内容审核范围: 封面信息、指导教师姓名、封面填写是否完整等封面事项一律不审核、不评价,
   评语任何部分(总体评价/具体评价/修改意见)均不得提及封面问题。
6. 其余分项字段照常输出(供内部判定), 但"评审报告"是最终呈现文本, 必须自足完整。

请严格按以下JSON格式输出详细的评审结果。每个评语都要详细具体，指出问题所在。

        {{
            "是否合格": true/false,
  "评审报告": "【总体评价】...【具体评价】...【修改意见】...",
            "总分": 60-100之间的整数,
            "各项评分": {{
                "题目": {{
                    "得分": 0-15,
                    "评语": "详细的题目评价，指出是否符合要求，如符合请说明包含的具体研究对象",
                    "问题": ["具体问题1", "具体问题2"]
                }},
                "字数": {{
                    "得分": 0-10,
                    "评语": "详细的字数评价，说明实际字数与要求的差距",
                    "问题": ["具体问题1"]
                }},
                "结构": {{
                    "得分": 0-25,
                    "评语": "详细的结构评价，说明各部分是否完整，逻辑是否正确，是否有重复论述",
                    "包含的章节": ["章节1", "章节2", "..."],
                    "缺失的章节": ["缺失部分1", "缺失部分2"],
                    "是否有重复论述": true/false,
                    "重复论述说明": "如果有重复论述，请详细说明",
                    "问题": ["具体问题1", "具体问题2"]
                }},
                "语言逻辑": {{
                    "得分": 0-20,
                    "评语": "详细的语言评价，说明语言是否专业通顺，逻辑是否合理，理论应用是否得当",
                    "问题": ["具体问题1", "具体问题2"]
                }},
                "数据支撑": {{
                    "得分": 0-20,
                    "评语": "详细的数据支撑评价，说明图表数量和质量，是否充分支撑分析",
                    "图表统计": {{
                        "表格数量": 0,
                        "图片数量": 0,
                        "图表质量": "评价图表的质量和相关性"
                    }},
                    "问题": ["具体问题1", "具体问题2"]
                }},
                "格式": {{
                    "得分": 0-10,
                    "评语": "详细的格式评价，说明格式是否符合要求，上下文是否一致",
                    "问题": ["具体问题1", "具体问题2"]
                }}
            }},
            "必达项检查": {{
                "题目合格": true/false,
                "字数合格": true/false,
                "结构合格": true/false,
                "语言合格": true/false,
                "数据合格": true/false
            }},
            "总体评语": "详细的总体评价，总结论文的优缺点，指出主要问题",
            "修改建议": [
                "详细的修改建议1",
                "详细的修改建议2",
                "详细的修改建议3",
                "详细的修改建议4"
            ]
        }}
        """
        
        messages = [
            {"role": "system", "content": "你是一个严格的本科毕业论文评审专家，必须严格按照学校要求评审。1/2/3/4/5必须全部达到，否则不合格。请给出详细的评语和建议。"},
            {"role": "user", "content": prompt}
        ]
        
        # 多轮完整调用: 每次解析失败都重新调用API(LLM偶发输出非法JSON, 重调通常成功),
        # 而非对同一份content空转重试; 清理markdown围栏后 json.loads, 失败再整段提取
        last_err, content = None, ""
        for attempt in range(3):
            if attempt == 0:
                resp = self.call_deepseek_api(messages, max_tokens=4000)
            else:
                self.logger.warning(f"JSON解析失败, 重新调用API(第{attempt+1}/3): {last_err}")
                resp = self.call_deepseek_api(messages, max_tokens=8000)
            if not (resp and "choices" in resp):
                last_err = "API调用无返回"
                time.sleep(2)
                continue
            content = resp["choices"][0]["message"].get("content", "") or ""
            if not content:
                last_err = "空正文"
                continue
            try:
                cleaned = content.strip()
                if cleaned.startswith("```"):
                    cleaned = re.sub(r"^```[a-zA-Z]*\s*|\s*```$", "", cleaned, flags=re.S)
                i, j = cleaned.find("{"), cleaned.rfind("}")
                if i >= 0 and j > i:
                    cleaned = cleaned[i:j + 1]
                return json.loads(cleaned)
            except Exception as e:
                last_err = e
                self.logger.error(f"JSON解析失败(第{attempt+1}/3): {e}")
            time.sleep(1)
        self.logger.warning("LLM输出无法解析为JSON, 改用规则兜底")
        # API调用失败，使用规则审核
        return self._rule_based_review(thesis_text, thesis_info, title)
    
    def _rule_based_review(self, thesis_text: str, thesis_info: Dict, title: str) -> Dict:
        """
        基于规则的备用审核方案（返回详细评语）
        """
        stats = thesis_info.get("stats", {})
        sections = thesis_info.get("sections", {})
        
        # 初始化详细的评分和评语
        scores = {
            "题目": {"得分": 10, "评语": "", "问题": []},
            "字数": {"得分": 5, "评语": "", "问题": []},
            "结构": {"得分": 15, "评语": "", "包含的章节": [], "缺失的章节": [], "是否有重复论述": False, "重复论述说明": "", "问题": []},
            "语言逻辑": {"得分": 12, "评语": "", "问题": []},
            "数据支撑": {"得分": 10, "评语": "", "图表统计": {"表格数量": 0, "图片数量": 0, "图表质量": ""}, "问题": []},
            "格式": {"得分": 5, "评语": "", "问题": []}
        }
        
        issues = []
        suggestions = []
        
        # 1. 题目检查
        title_qualified, title_msg = self.check_title_format(title)
        if title_qualified:
            scores["题目"]["得分"] = 15
            scores["题目"]["评语"] = f"题目符合要求，包含具体研究对象。{title_msg}"
        else:
            scores["题目"]["得分"] = 5
            scores["题目"]["问题"].append(title_msg)
            issues.append(f"题目问题：{title_msg}")
            suggestions.append("请修改题目，增加具体的企业/项目/案例名称")
        
        # 2. 字数检查
        word_count = stats.get("word_count", 0)
        if word_count >= 10000:
            scores["字数"]["得分"] = 10
            scores["字数"]["评语"] = f"字数达标：{word_count}字，符合10000字以上的要求。"
        else:
            scores["字数"]["得分"] = max(0, 10 - int((10000 - word_count) / 1000))
            shortage = 10000 - word_count
            scores["字数"]["问题"].append(f"字数不足：当前{word_count}字，要求10000字，还差{shortage}字")
            issues.append(f"字数不足：当前{word_count}字")
            suggestions.append(f"必须扩充内容，还需增加{shortage}字")
        
        # 3. 结构检查
        section_keywords = {
            "引言": ["引言", "绪论"],
            "理论基础": ["理论基础", "研究现状", "文献综述"],
            "现状": ["现状"],
            "问题分析": ["问题", "原因"],
            "解决方案": ["解决方案", "策略", "对策"],
            "结论": ["结论"]
        }
        
        found_sections = []
        missing = []
        
        for part, keywords in section_keywords.items():
            found = False
            for section in sections.keys():
                for keyword in keywords:
                    if keyword in section:
                        found = True
                        found_sections.append(part)
                        break
                if found:
                    break
            if not found:
                missing.append(part)
        
        if not missing:
            scores["结构"]["得分"] = 25
            scores["结构"]["评语"] = "结构完整，包含所有必要部分，逻辑清晰。"
            scores["结构"]["包含的章节"] = found_sections
        else:
            scores["结构"]["得分"] = max(0, 25 - len(missing) * 4)
            scores["结构"]["缺失的章节"] = missing
            scores["结构"]["评语"] = f"结构不完整，缺少以下部分：{', '.join(missing)}"
            scores["结构"]["问题"].append(f"缺少必要内容：{', '.join(missing)}")
            issues.append(f"结构不完整，缺少：{', '.join(missing)}")
            suggestions.append(f"必须补充以下内容：{', '.join(missing)}")
        
        # 检查重复论述（简单规则判断）
        # 这里可以加入更复杂的重复检测逻辑
        scores["结构"]["是否有重复论述"] = False
        
        # 4. 数据支撑检查
        table_count = stats.get("table_count", 0)
        image_count = stats.get("image_count", 0)
        scores["数据支撑"]["图表统计"]["表格数量"] = table_count
        scores["数据支撑"]["图表统计"]["图片数量"] = image_count
        
        if table_count >= 2 or image_count >= 2:
            scores["数据支撑"]["得分"] = 20
            scores["数据支撑"]["评语"] = f"数据支撑充分，有{table_count}个表格，{image_count}张图片，能够有效支撑分析。"
            scores["数据支撑"]["图表统计"]["图表质量"] = "图表数量充足，能够有效支撑分析"
        elif table_count >= 1 or image_count >= 1:
            scores["数据支撑"]["得分"] = 15
            scores["数据支撑"]["评语"] = f"数据支撑基本够用，有{table_count}个表格，{image_count}张图片，但可以进一步丰富。"
            scores["数据支撑"]["图表统计"]["图表质量"] = "图表数量基本够用"
            scores["数据支撑"]["问题"].append("数据支撑略显不足")
            issues.append("数据支撑略显不足")
            suggestions.append("建议增加更多数据表格或图表")
        else:
            scores["数据支撑"]["得分"] = 5
            scores["数据支撑"]["评语"] = "完全没有数据或图表支撑，不符合要求。"
            scores["数据支撑"]["问题"].append("完全没有数据或图表支撑")
            issues.append("缺乏数据支撑")
            suggestions.append("必须增加数据表格或图表进行分析支撑")
        
        # 5. 语言逻辑（规则判断较难，给中等分）
        scores["语言逻辑"]["得分"] = 15
        scores["语言逻辑"]["评语"] = "语言基本通顺，逻辑基本合理，但需要大模型进行更准确的判断。"
        
        # 6. 格式（规则判断较难，给中等分）
        scores["格式"]["得分"] = 7
        scores["格式"]["评语"] = "格式基本符合要求，上下文格式基本一致。封面或标题页格式问题不做严格要求。"
        
        # 7. 计算总分
        total_score = sum([item["得分"] for item in scores.values()])
        
        # 8. 必达项检查
        must_pass = {
            "题目合格": title_qualified,
            "字数合格": word_count >= 10000,
            "结构合格": len(missing) == 0,
            "语言合格": True,  # 规则判断难以检查语言，默认为合格
            "数据合格": not (table_count == 0 and image_count == 0)
        }
        
        is_qualified = all(must_pass.values())
        
        # 9. 总体评语
        if is_qualified:
            overall = f"论文达到基本要求。"
        else:
            overall = f"论文未达到基本要求。"
        
        if issues:
            overall += f"主要问题：{'; '.join(issues[:3])}"
        
        return {
            "是否合格": is_qualified,
            "总分": total_score,
            "各项评分": scores,
            "必达项检查": must_pass,
            "总体评语": overall,
            "修改建议": suggestions if suggestions else ["论文基本符合要求，可进一步优化"]
        }
    
    def generate_report(self, review_result: Dict, thesis_info: Dict, title: str) -> str:
        """
        生成详细的审核报告
        """
        report = []
        report.append("="*70)
        report.append("无锡太湖学院本科毕业论文审核报告")
        report.append("="*70)
        report.append("")
        
        # 基本信息
        report.append("一、论文基本信息")
        report.append("-"*50)
        report.append(f"论文题目：{title}")
        report.append(f"论文文件：{thesis_info.get('file_name', '未知')}")
        report.append(f"总字数：{thesis_info.get('stats', {}).get('word_count', 0)}字")
        report.append(f"表格数量：{thesis_info.get('stats', {}).get('table_count', 0)}个")
        report.append(f"图片数量：{thesis_info.get('stats', {}).get('image_count', 0)}张")
        report.append("")
        
        # 审核结论
        report.append("二、审核结论")
        report.append("-"*50)
        report.append(f"审核结果：{'合 格' if review_result.get('是否合格') else '不合格'}")
        report.append(f"最终得分：{review_result.get('总分', 0)}分")
        report.append("")
        
        # 必达项检查（1/2/3/4/5必须达到，否则不合格）
        report.append("三、必达项检查（第1/2/3/4/5条必须全部达到，否则不合格）")
        report.append("-"*50)
        must_pass = review_result.get('必达项检查', {})
        if must_pass:
            items = [
                ("题目合格", "1. 题目以具体企业/项目/案例为研究对象"),
                ("字数合格", "2. 字数达到10000字以上"),
                ("结构合格", "3. 包含六部分基本结构（标题可略有不同，逻辑正确）"),
                ("语言合格", "4. 语言专业通顺，逻辑合理，能应用理论分析问题"),
                ("数据合格", "5. 有适当的数据或图表作为分析支撑")
            ]
            for key, desc in items:
                result = must_pass.get(key, False)
                report.append(f"   {'✓' if result else '✗'} {desc}")
        report.append("")
        
        # 分项评分（详细评语）
        report.append("四、分项评分与详细评语")
        report.append("-"*50)
        scores = review_result.get('各项评分', {})
        
        for item, score_info in scores.items():
            report.append(f"\n【{item}】{score_info.get('得分', 0)}分")
            
            if score_info.get('评语'):
                report.append(f"评语：{score_info['评语']}")
            
            # 结构部分的详细信息
            if item == "结构":
                if score_info.get('包含的章节'):
                    report.append(f"包含的章节：{', '.join(score_info['包含的章节'])}")
                if score_info.get('缺失的章节'):
                    report.append(f"缺失的章节：{', '.join(score_info['缺失的章节'])}")
                if score_info.get('是否有重复论述'):
                    report.append(f"重复论述：{'是' if score_info['是否有重复论述'] else '否'}")
                    if score_info.get('重复论述说明'):
                        report.append(f"重复论述说明：{score_info['重复论述说明']}")
            
            # 数据支撑部分的详细信息
            if item == "数据支撑" and score_info.get('图表统计'):
                charts = score_info['图表统计']
                report.append(f"图表统计：{charts.get('表格数量', 0)}个表格，{charts.get('图片数量', 0)}张图片")
                if charts.get('图表质量'):
                    report.append(f"图表质量：{charts['图表质量']}")
            
            # 问题列表
            if score_info.get('问题'):
                report.append("存在问题：")
                for issue in score_info['问题']:
                    report.append(f"  • {issue}")
        
        report.append("")
        
        # 总体评语
        report.append("五、总体评语")
        report.append("-"*50)
        report.append(review_result.get('总体评语', '无'))
        report.append("")
        
        # 修改建议
        report.append("六、修改建议")
        report.append("-"*50)
        suggestions = review_result.get('修改建议', [])
        if suggestions:
            for i, suggestion in enumerate(suggestions, 1):
                report.append(f"{i}. {suggestion}")
        else:
            report.append("论文符合要求，无需修改。")
        
        report.append("")
        report.append("="*70)
        report.append(f"审核时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        report.append("="*70)
        
        return "\n".join(report)
    
    def save_report(self, report: str, output_path: str = None):
        """保存审核报告"""
        if not output_path:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            output_path = f"论文审核报告_{timestamp}.txt"
        
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(report)
        
        self.logger.info(f"审核报告已保存至：{output_path}")
        return output_path
    
    def review_thesis(self, file_path: str) -> Dict:
        """
        审核单篇论文
        """
        self.logger.info(f"开始审核论文：{file_path}")
        
        # 1. 提取论文内容
        thesis_text, thesis_info = self.extract_text_from_docx(file_path)
        if not thesis_text:
            self.logger.error("论文内容提取失败")
            return None
        
        # 2. 提取题目
        doc = Document(file_path)
        title = self.extract_title_from_doc(doc)
        
        self.logger.info(f"论文题目：{title}")
        self.logger.info(f"字数：{thesis_info['stats']['word_count']}字，表格：{thesis_info['stats']['table_count']}个")
        
        # 3. 使用DeepSeek审核
        self.logger.info("正在使用DeepSeek进行审核...")
        thesis_info['_source_path'] = file_path
        review_result = self.review_with_deepseek(thesis_text, thesis_info, title)
        
        # 4. 生成报告
        report = self.generate_report(review_result, thesis_info, title)
        
        # 5. 保存报告
        self.save_report(report)
        
        # 6. 打印结果
        self.logger.info(f"审核完成！")
        self.logger.info(f"结果：{'合格' if review_result.get('是否合格') else '不合格'}")
        self.logger.info(f"得分：{review_result.get('总分')}分")
        
        return {
            "result": review_result,
            "report": report,
            "thesis_info": thesis_info
        }


def main():
    """
    主函数
    """
    print("="*70)
    print("无锡太湖学院本科毕业论文审核程序")
    print("（使用DeepSeek大模型）")
    print("="*70)
    print()
    print("【审核说明】")
    print("1. 题目必须包含具体企业/项目/案例名称")
    print("2. 字数必须达到10000字以上")
    print("3. 必须包含六部分基本结构（标题可略有不同）")
    print("4. 语言必须专业通顺，逻辑合理")
    print("5. 必须有数据或图表支撑")
    print("="*70)
    print()
    
    # 使用提供的API密钥
    api_key = os.environ.get("DEEPSEEK_API_KEY", "")
    
    # 获取论文文件路径
    while True:
        file_path = input("请输入论文文件路径（支持.docx格式）：").strip()
        if not file_path:
            print("错误：必须指定论文文件路径")
            continue
        
        if not os.path.exists(file_path):
            print(f"错误：文件不存在 - {file_path}")
            continue
        
        if not file_path.endswith('.docx'):
            print("警告：程序主要支持.docx格式，其他格式可能无法正确读取")
            confirm = input("是否继续？(y/n，默认n)：").strip().lower()
            if confirm != 'y':
                continue
        
        break
    
    print("\n正在初始化审核程序...")
    
    # 创建审核器
    reviewer = ThesisDeepSeekReviewer(api_key=api_key)
    
    # 执行审核
    print("\n开始审核，请稍候（约30-60秒）...")
    result = reviewer.review_thesis(file_path)
    
    if result:
        print("\n" + "="*70)
        print("审核完成！")
        print(f"审核结果：{'合 格' if result['result'].get('是否合格') else '不合格'}")
        print(f"最终得分：{result['result'].get('总分')}分")
        print("\n详细报告已保存到文件：论文审核报告_时间戳.txt")
        print("="*70)
    else:
        print("\n审核失败，请检查日志。")


if __name__ == "__main__":
    main()