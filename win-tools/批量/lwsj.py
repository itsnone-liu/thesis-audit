"""
无锡太湖学院设计类本科毕业论文审核程序
文件名：lwsj.py
功能：使用DeepSeek大模型按照设计类论文要求评审论文
"""

import os
import json
import re
import time
from docx import Document
from datetime import datetime
import requests
from typing import Dict, List, Tuple, Optional
import logging



# ---- 新口径同步 (2026-09-21): 四形态对象认定 + 审计条款A-F + 确定性预检 ----
import os as _os, sys as _sys
_sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", "增强0908"))
try:
    from audit_principles import PRINCIPLES_TEXT, precheck as _precheck, evidence_block as _evidence_block
except Exception:
    PRINCIPLES_TEXT, _precheck, _evidence_block = "", None, None

OBJ_RULE = ("【研究对象认定(2026-09-20口径)】研究对象四形态均有效: ①具体名称前置 ②脱敏表述前置"
            "('某高校''某企业''某公司''A公司''B集团'等) ③具体名称副标题'——以XX为例' ④脱敏副标题"
            "'——以某XX为例'。只判定是否存在可锚定的研究对象, 不得因脱敏/不具名/位置形态扣分或判不合格。")

def _pa_block(file_path):
    """审计条款+实测证据prompt块(模块未加载时返回空串, 不影响原流程)。"""
    if _precheck is None or not file_path:
        return ""
    try:
        return PRINCIPLES_TEXT + "\n## 实测证据(程序预检, 优先于目测)\n" + _evidence_block(_precheck(file_path))
    except Exception:
        return ""

class ThesisDesignReviewer:
    """
    使用DeepSeek的设计类论文审核程序
    适用于建筑、环境设计、艺术设计等专业的论文评审
    """
    
    def __init__(self, api_key: str = None):
        """
        初始化审核器
        :param api_key: DeepSeek API密钥
        """
        self.logger = self._setup_logger()
        self.api_key = api_key or os.environ.get("DEEPSEEK_API_KEY", "")  # 2026-09-21修: 读env(百炼), 原硬编码官方key+百炼端点=401
        self.api_base = os.environ.get("DEEPSEEK_BASE_URL", "https://api.deepseek.com/v1")
        self.model = os.environ.get("DEEPSEEK_MODEL", "deepseek-chat")
        
        if not self.api_key:
            self.logger.warning("未设置API密钥，将使用规则进行基础审核")
        
        # 设计类论文审核要求
        self.requirements = {
            "题目要求": "有具体的设计对象，以具体建筑、物体、环境氛围为研究对象(脱敏表述同四形态有效, 见OBJ_RULE)",
            "字数要求": 10000,
            "结构要求": {
                "必须包含的部分": [
                    "第一部分：引言（绪论）",
                    "第二部分：相关设计理论及背景",
                    "第三部分：问题发现与分析",
                    "第四部分：设计策略",
                    "第五部分：设计方案（需配图纸）",
                    "第六部分：总结与展望"
                ],
                "注意事项": [
                    "每章的标题可以表述上略有不同",
                    "设计策略与设计方案可以合为一章"
                ],
                "图纸要求": "需配1-2张设计图（平面图/立面图/剖面图）和至少一张效果图"
            },
            "语言要求": "语句通顺专业，不能有口语化表达，逻辑合理，能够应用理论完成设计",
            "格式要求": "格式基本符合要求，文章本身的上下文格式必须保持一致（封面或标题页格式问题不做要求）"
        }
        
        # 评分标准
        self.scoring_criteria = {
            "题目": {"weight": 15, "description": "是否有具体设计对象"},
            "字数": {"weight": 10, "description": "是否达到10000字以上"},
            "结构框架": {"weight": 25, "description": "是否包含六部分，逻辑是否正确（问题→策略→方案）"},
            "理论应用": {"weight": 15, "description": "是否应用理论，理论与设计是否关联"},
            "设计策略": {"weight": 15, "description": "策略是否合理，是否回应问题"},
            "设计方案": {"weight": 20, "description": "方案是否具体，图纸是否满足要求，策略与方案是否对应"}
        }
    
    def _setup_logger(self) -> logging.Logger:
        """设置日志记录器"""
        logger = logging.getLogger("DesignReviewer")
        logger.setLevel(logging.INFO)
        
        ch = logging.StreamHandler()
        ch.setLevel(logging.INFO)
        formatter = logging.Formatter('%(asctime)s - %(levelname)s - %(message)s')
        ch.setFormatter(formatter)
        logger.addHandler(ch)
        
        return logger
    
    def call_deepseek_api(self, messages: List[Dict], temperature: float = 0.3, max_tokens: int = 5000) -> Dict:
        """调用DeepSeek API"""
        headers = {
            "Authorization": f"Bearer {self.api_key}",
            "Content-Type": "application/json"
        }
        data = {
            "model": self.model,
            "enable_thinking": False,
            "messages": messages,
            "temperature": temperature,
            "max_tokens": max_tokens,
            "response_format": {"type": "json_object"}
        }
        
        import time as _time
        for _attempt in range(3):
            try:
                response = requests.post(
                    f"{self.api_base}/chat/completions",
                    headers=headers,
                    json=data,
                    timeout=90
                )
                if response.status_code == 200:
                    return response.json()
                self.logger.error(f"API调用失败(第{_attempt+1}次)：{response.status_code} - {response.text[:120]}")
            except Exception as e:
                self.logger.error(f"API调用异常(第{_attempt+1}次)：{str(e)[:120]}")
            if _attempt < 2:
                _time.sleep(2 * (_attempt + 1))
        return None
    
    def extract_text_from_docx(self, file_path: str) -> Tuple[str, Dict]:
        """从docx文件中提取文本和结构信息"""
        try:
            doc = Document(file_path)
            full_text = []
            sections = {}
            current_section = "开头"
            section_content = []
            
            stats = {
                "paragraph_count": 0,
                "table_count": len(doc.tables),
                "image_count": 0,
                "word_count": 0,
                "design_images": [],      # 记录图片信息
                "design_image_count": 0,  # 设计图数量（平面图/立面图/剖面图）
                "effect_image_count": 0,  # 效果图数量（效果图/渲染图）
                "unlabeled_images": []    # 未标注的图片
            }
            
            # 第一遍：收集所有图片位置和周围文字
            image_index = 0
            images_info = []  # 存储每个图片的段落索引和周围文字
            
            for para_idx, para in enumerate(doc.paragraphs):
                text = para.text.strip()
                # 2026-09-22修: VML老格式(<w:pict>/<v:imagedata>)也认, 不再只认graphicData
                para_xml = para._element.xml
                # 2026-09-22修: 段内多图也逐张计数(同一段放多张图/拼图时, 原逻辑每段只+1会漏数)
                #   <a:blip>计数精确(每张图一个r:embed), VML老格式数<v:imagedata>
                n_imgs = para_xml.count('<a:blip') + para_xml.count('<v:imagedata')
                if ('graphicData' in para_xml or '<w:pict' in para_xml
                        or 'imagedata' in para_xml or '<w:drawing' in para_xml) and n_imgs > 0:
                    # 获取图片周围的文字（上一段 + 当前段落 + 下一段落）
                    # 2026-09-22修: 图注("图 4-1 xxx示意图")可能位于图片上方或下方, 三向合并
                    surrounding_text = ""
                    if para_idx > 0:
                        surrounding_text += doc.paragraphs[para_idx - 1].text.strip().lower() + " "
                    surrounding_text += text.lower()
                    if para_idx + 1 < len(doc.paragraphs):
                        next_text = doc.paragraphs[para_idx + 1].text.strip().lower()
                        surrounding_text += " " + next_text
                    for _ in range(n_imgs):
                        image_index += 1
                        images_info.append({
                            "index": image_index,
                            "para_idx": para_idx,
                            "surrounding_text": surrounding_text
                        })
                        stats["image_count"] += 1
            
            # 2026-09-22修: 扫描表格中的图片——学生常把设计图/效果图放进表格单元格,
            # 只遍历doc.paragraphs会漏数, 导致"图纸严重不足"误判。图注一般在同一cell内文字。
            for tb_idx, tb in enumerate(doc.tables):
                # 整表文字兜底(表头列名常含"设计图/效果图/示意图"等)
                table_all_text = " ".join(
                    cell.text.strip().lower()
                    for row in tb.rows for cell in row.cells if cell.text.strip()
                )
                for row in tb.rows:
                    for cell in row.cells:
                        cell_xml = cell._tc.xml
                        n_imgs = cell_xml.count('<a:blip') + cell_xml.count('<v:imagedata')
                        if n_imgs > 0:
                            cell_text = cell.text.strip().lower()
                            surrounding_text = (cell_text + " " + table_all_text).strip()
                            for _ in range(n_imgs):
                                image_index += 1
                                images_info.append({
                                    "index": image_index,
                                    "para_idx": -1,
                                    "in_table": True,
                                    "tb_idx": tb_idx,
                                    "surrounding_text": surrounding_text
                                })
                                stats["image_count"] += 1
            
            # 第二遍：根据周围文字判断图片类型
            # 2026-09-22修: 扩展关键词——设计类论文图注常用"示意图/系统图/应用图/展示图/设计图",
            # 效果图常用"效果图/渲染图/场景/展示/视觉呈现", 覆盖视觉传达/环境设计/服装等多专业
            design_keywords = ['平面图', '立面图', '剖面图', '设计图', '总平面', '布局图', '施工图',
                               '示意图', '系统图', '应用图', '展示图', '设计稿', 'logo', '标志',
                               'vi', '视觉识别', '图形设计', '图案', '纹样', '款式图', '结构图']
            effect_keywords = ['效果图', '渲染图', '透视图', '三维图', '3d', '鸟瞰图',
                               '场景应用', '场景展示', '场景呈现', '应用场景', '展示效果',
                               '视觉呈现', '陈列', '实景', '虚拟展示']
            
            for img in images_info:
                text = img["surrounding_text"]
                is_design = any(kw in text for kw in design_keywords)
                is_effect = any(kw in text for kw in effect_keywords)
                
                if is_design:
                    stats["design_images"].append({"type": "设计图", "index": img["index"]})
                    stats["design_image_count"] += 1
                elif is_effect:
                    stats["design_images"].append({"type": "效果图", "index": img["index"]})
                    stats["effect_image_count"] += 1
                else:
                    stats["design_images"].append({"type": "未识别", "index": img["index"]})
                    stats["unlabeled_images"].append(img["index"])
            
            # 提取章节内容
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    full_text.append(text)
                    stats["paragraph_count"] += 1
                    
                    if self._is_section_title(text):
                        if section_content:
                            sections[current_section] = "\n".join(section_content)
                        current_section = text
                        section_content = []
                    else:
                        section_content.append(text)
            
            if section_content:
                sections[current_section] = "\n".join(section_content)
            
            full_text_str = "\n".join(full_text)
            stats["word_count"] = len(full_text_str)
            
            return full_text_str, {
                "sections": sections,
                "stats": stats,
                "file_name": os.path.basename(file_path)
            }
            
        except Exception as e:
            self.logger.error(f"读取文档失败：{str(e)}")
            return "", {}
    
    def _is_section_title(self, text: str) -> bool:
        """判断是否为章节标题"""
        # 2026-09-21修: ①目录行(标题后跟\t页码/多空格页码)不当标题 ②支持"第1章"阿拉伯数字
        if re.search(r'\t\s*\d{1,3}$', text) or re.search(r'\s{2,}\d{1,3}$', text):
            return False
        patterns = [
            r'^第[一二三四五六七八九十\d]+章',
            r'^摘要$|^abstract$|^目录$|^绪论$|^引言$|^结论$|^参考文献$|^致谢$',
            r'^\d+\.\d+\s+',
            r'^第一部分|^第二部分|^第三部分|^第四部分|^第五部分|^第六部分'
        ]
        for pattern in patterns:
            if re.search(pattern, text, re.IGNORECASE):
                return True
        return False
    
    def check_title_format(self, title: str) -> Tuple[bool, str]:
        """检查题目格式：有具体设计对象"""
        patterns = [
            r'.*以.*为例', r'.*设计研究.*', r'.*方案研究.*',
            r'.*公园.*', r'.*咖啡馆.*', r'.*小区.*', r'.*建筑.*',
            r'.*室内.*', r'.*景观.*', r'.*环境.*', r'.*改造.*'
        ]
        for pattern in patterns:
            if re.search(pattern, title):
                return True, "题目包含具体设计对象"
        return False, "题目缺少具体设计对象（必须包含具体建筑/物体/环境氛围）"
    
    def extract_title_from_doc(self, doc) -> str:
        """从文档中提取题目(2026-09-21修: 优先解析'论文题目：xxx'标签行, 修复误抓封面校名行)"""
        try:
            # ① 显式标签行: "论文题目：xxx" / "题 目：xxx"(含副标题全长)
            for para in doc.paragraphs[:40]:
                t = para.text.strip()
                m = re.match(r"^论文题目[:：]\s*(.+)$", t) or re.match(r"^题\s*目[:：]\s*(.+)$", t)
                if m and m.group(1).strip():
                    return m.group(1).strip()
            # ② 启发式fallback: 排除校名/学位/答辩类封面行
            skip = re.compile(r"学院|大学|自学考试|本科|专科|专升本|答辩|专业|姓名|学号|指导|日期|年|月")
            for para in doc.paragraphs[:30]:
                t = para.text.strip()
                if 12 <= len(t) <= 60 and not skip.search(t) and "摘" not in t:
                    return t
        except:
            pass
        return "未找到题目"
    
    def check_image_requirements(self, stats: Dict) -> Tuple[bool, str, Dict]:
        """
        检查图片要求：
        - 设计图（平面图/立面图/剖面图）至少1张
        - 效果图（效果图/渲染图）至少1张
        - 如果图片数量达标但无图注，提醒添加
        """
        design_count = stats.get("design_image_count", 0)
        effect_count = stats.get("effect_image_count", 0)
        unlabeled = stats.get("unlabeled_images", [])
        
        detail = {
            "设计图数量": design_count,
            "效果图数量": effect_count,
            "总图片数": stats.get("image_count", 0),
            "未识别图片": unlabeled
        }
        
        # 判断图纸是否达标
        has_design = design_count >= 1
        has_effect = effect_count >= 1
        qualified = has_design and has_effect
        
        # 生成评语
        if qualified:
            if unlabeled:
                message = f"图纸数量达标（设计图{design_count}张，效果图{effect_count}张），但{len(unlabeled)}张图片未标注图注"
            else:
                message = f"图纸符合要求：设计图{design_count}张，效果图{effect_count}张"
        else:
            issues = []
            if not has_design:
                issues.append(f"缺少设计图（需要平面图/立面图/剖面图）")
            if not has_effect:
                issues.append(f"缺少效果图（需要效果图/渲染图）")
            message = f"图纸不符合要求：{'；'.join(issues)}"
        
        return qualified, message, detail
    
    def review_with_deepseek(self, thesis_text: str, thesis_info: Dict, title: str) -> Dict:
        """使用DeepSeek审核设计类论文"""
        stats = thesis_info.get("stats", {})
        image_qualified, image_message, image_detail = self.check_image_requirements(stats)
        
        _pa = _pa_block(thesis_info.get('_source_path'))
        prompt = f"""
        你是一位严格的本科设计类毕业论文评审专家。请按照以下要求对论文进行详细评审。

        ## 论文基本信息
        - 题目：{title}
        - 总字数：{stats.get('word_count', 0)}字
        - 图纸情况：设计图{image_detail.get('设计图数量', 0)}张，效果图{image_detail.get('效果图数量', 0)}张
        - 未标注图片：{len(image_detail.get('未识别图片', []))}张

        ## 审核要求

        ### 必达项（1/2/4必须达到，否则不合格；3要求框架逻辑正确+图纸必须有）
        1. 题目：有具体的设计对象（以具体建筑、物体、环境氛围为研究对象）
           {OBJ_RULE}
        2. 字数：10000字以上
        3. 框架逻辑：大框架逻辑必须正确（发现问题→提出策略→设计方案）
           图纸：必须有（设计图+效果图，数量不论，有即可）
        4. 语言：语句通顺专业，不能有口语化表达，逻辑合理

        ### 结构要求
        必须包含以下部分（标题可以略有不同）：
        - 第一部分：引言（绪论）
        - 第二部分：相关设计理论及背景
        - 第三部分：问题发现与分析
        - 第四部分：设计策略
        - 第五部分：设计方案（总体布局、分区设计、专项设计）
        - 第六部分：总结与展望

        ## 论文内容（前6000字）
        {thesis_text[:6000]}

        {_pa}

        
## 评审报告撰写规范(最高优先级, 适用于所有评语与"评审报告"字段)
1. 评语以自然语言撰写, 像一位毕业论文指导教师在评阅: 禁止"实测""检测到""程序判定""证据""预检""指标显示"等程序化词汇;
   数据表述用自然形式(如"全文约1.4万字, 插图14幅、表格3张")。
2. JSON新增顶层字段 "评审报告"(字符串): 一篇完整评语, 严格按【总体评价】→【具体评价】→【修改意见】三段固定顺序组织,
   三段先后顺序不得颠倒, 且【总体评价】【具体评价】【修改意见】三个标记必须各自另起一行独占段首,
   段与段之间空一行, 任何标记不得接在上一段末尾同一行。
【总体评价】2-4句: 论文整体质量与结论。若合格, 写明建议成绩(如"综合评定为合格, 建议成绩82分");
  若不合格, 只作定性结论, 不得出现任何分数、分值、得分。
3. 【具体评价】(在"评审报告"内): 按题目、结构、内容与论证、语言与规范等维度逐段展开,
   所有具体问题在此详细指出(引到章节、数值、图表编号等原文位置), 用连贯段落, 不用表格罗列。
4. 【修改意见】(在"评审报告"内): 逐条列出, 具体可执行; 不合格论文的修改意见要覆盖全部主要缺陷。
5. 封面页不属于论文内容审核范围: 封面信息、指导教师姓名、封面填写是否完整等封面事项一律不审核、不评价,
   评语任何部分(总体评价/具体评价/修改意见)均不得提及封面问题。
6. 其余分项字段照常输出(供内部判定), 但"评审报告"是最终呈现文本, 必须自足完整。

请按以下JSON格式输出评审结果：

        {{
            "是否合格": true/false,
  "评审报告": "【总体评价】...【具体评价】...【修改意见】...",
            "总分": 60-100之间的整数,
            "必达项检查": {{
                "题目合格": true/false,
                "字数合格": true/false,
                "框架逻辑合格": true/false,
                "图纸合格": true/false,
                "语言合格": true/false
            }},
            "各项评分": {{
                "题目": {{"得分": 0-15, "评语": "", "问题": []}},
                "字数": {{"得分": 0-10, "评语": "", "问题": []}},
                "结构框架": {{"得分": 0-25, "评语": "", "包含的章节": [], "缺失的章节": [], "逻辑问题": []}},
                "理论应用": {{"得分": 0-15, "评语": "", "问题": []}},
                "设计策略": {{"得分": 0-15, "评语": "", "问题": []}},
                "设计方案": {{"得分": 0-20, "评语": "", "问题": []}}
            }},
            "总体评语": "",
            "修改建议": []
        }}
        """
        
        messages = [
            {"role": "system", "content": "你是严格的设计类本科毕业论文评审专家。第1/2/4条必须达到，第3条框架逻辑必须正确且图纸必须有，否则不合格。"},
            {"role": "user", "content": prompt}
        ]
        
        # 多轮完整调用: 每次解析失败都重新调用API(LLM偶发输出非法JSON, 重调通常成功),
        # 而非对同一份content空转重试; 清理markdown围栏后 json.loads, 失败再整段提取
        last_err, content = None, ""
        for attempt in range(3):
            if attempt == 0:
                resp = self.call_deepseek_api(messages, max_tokens=5000)
            else:
                self.logger.warning(f"JSON解析失败, 重新调用API(第{attempt+1}/3): {last_err}")
                resp = self.call_deepseek_api(messages, max_tokens=8000)
            if not (resp and "choices" in resp):
                last_err = "API调用无返回"
                time.sleep(2)
                continue
            content = resp["choices"][0]["message"].get("content", "") or ""
            if not content:
                last_err = "空正文"
                continue
            try:
                cleaned = content.strip()
                if cleaned.startswith("```"):
                    cleaned = re.sub(r"^```[a-zA-Z]*\s*|\s*```$", "", cleaned, flags=re.S)
                i, j = cleaned.find("{"), cleaned.rfind("}")
                if i >= 0 and j > i:
                    cleaned = cleaned[i:j + 1]
                return json.loads(cleaned)
            except Exception as e:
                last_err = e
                self.logger.error(f"JSON解析失败(第{attempt+1}/3): {e}")
            time.sleep(1)
        self.logger.warning("LLM输出无法解析为JSON, 改用规则兜底")
        return self._rule_based_review(thesis_text, thesis_info, title)
    
    def _rule_based_review(self, thesis_text: str, thesis_info: Dict, title: str) -> Dict:
        """基于规则的备用审核方案"""
        stats = thesis_info.get("stats", {})
        sections = thesis_info.get("sections", {})
        word_count = stats.get("word_count", 0)
        
        # 检查题目
        title_qualified, title_msg = self.check_title_format(title)
        
        # 检查字数
        word_qualified = word_count >= 10000
        
        # 检查图纸
        image_qualified, image_msg, image_detail = self.check_image_requirements(stats)
        
        # 检查结构
        required_keywords = ["引言", "绪论", "理论", "背景", "问题", "分析", "策略", "方案", "设计", "总结", "展望"]
        found_keywords = []
        for keyword in required_keywords:
            for section in sections.keys():
                if keyword in section:
                    found_keywords.append(keyword)
                    break
        structure_qualified = len(found_keywords) >= 5
        
        # 语言默认合格
        language_qualified = True
        
        # 必达项检查
        must_pass = {
            "题目合格": title_qualified,
            "字数合格": word_qualified,
            "框架逻辑合格": structure_qualified,
            "图纸合格": image_qualified,
            "语言合格": language_qualified
        }
        
        is_qualified = all(must_pass.values())
        
        # 计算得分
        total_score = 70
        if title_qualified:
            total_score += 5
        if word_qualified:
            total_score += 5
        if structure_qualified:
            total_score += 5
        if image_qualified:
            total_score += 10
        total_score = max(60, min(95, total_score))
        
        # 生成修改建议
        suggestions = self._generate_suggestions(must_pass, image_detail, thesis_info)
        
        return {
            "是否合格": is_qualified,
            "总分": total_score,
            "必达项检查": must_pass,
            "各项评分": {
                "题目": {"得分": 15 if title_qualified else 8, "评语": title_msg, "问题": []},
                "字数": {"得分": 10 if word_qualified else 5, "评语": f"当前字数：{word_count}字", "问题": []},
                "结构框架": {"得分": 20 if structure_qualified else 12, "评语": "结构基本完整" if structure_qualified else "结构不够完整", "包含的章节": list(sections.keys())[:5], "缺失的章节": [], "逻辑问题": []},
                "理论应用": {"得分": 12, "评语": "理论应用基本合理", "问题": []},
                "设计策略": {"得分": 12, "评语": "策略基本合理", "问题": []},
                "设计方案": {"得分": 16 if image_qualified else 10, "评语": image_msg, "问题": []}
            },
            "总体评语": f"论文{'达到' if is_qualified else '未达到'}基本要求。",
            "修改建议": suggestions
        }
    
    def _generate_suggestions(self, must_pass: Dict, image_detail: Dict, thesis_info: Dict) -> List[str]:
        """生成修改建议"""
        suggestions = []
        
        if not must_pass.get("题目合格"):
            suggestions.append("题目必须包含具体设计对象（建筑/物体/环境氛围），如'XX公园设计研究'、'XX咖啡馆室内设计'")
        
        if not must_pass.get("字数合格"):
            word_count = thesis_info.get('stats', {}).get('word_count', 0)
            need = 10000 - word_count
            suggestions.append(f"字数不足：当前{word_count}字，还需增加{need}字，建议扩充设计分析和方案阐述部分")
        
        if not must_pass.get("框架逻辑合格"):
            suggestions.append("框架逻辑不完整，请确保包含：引言、理论背景、问题分析、设计策略、设计方案、总结与展望")
        
        if not must_pass.get("图纸合格"):
            suggestions.append("图纸数量不足，请补充：1-2张设计图（平面图/立面图/剖面图）和至少1张效果图")
        else:
            # 图纸数量达标，检查是否有未标注的图片
            unlabeled = image_detail.get("未识别图片", [])
            if unlabeled:
                suggestions.append(f"检测到{len(unlabeled)}张图片未添加图注，请在每张图纸下方标注图注，格式如：'图4-1 平面图'、'图4-2 效果图'等")
            else:
                suggestions.append("图纸数量符合要求。请保持图注清晰，确保包含'平面图'、'效果图'等关键词")
        
        if not suggestions:
            suggestions.append("论文基本符合要求，可进一步优化论证深度和设计细节")
        
        return suggestions
    
    def generate_report(self, review_result: Dict, thesis_info: Dict, title: str) -> str:
        """生成详细的审核报告"""
        report = []
        report.append("="*70)
        report.append("无锡太湖学院设计类本科毕业论文审核报告")
        report.append("="*70)
        report.append("")
        
        # 基本信息
        report.append("一、论文基本信息")
        report.append("-"*50)
        report.append(f"论文题目：{title}")
        report.append(f"论文文件：{thesis_info.get('file_name', '未知')}")
        
        stats = thesis_info.get("stats", {})
        report.append(f"总字数：{stats.get('word_count', 0)}字")
        report.append(f"设计图：{stats.get('design_image_count', 0)}张")
        report.append(f"效果图：{stats.get('effect_image_count', 0)}张")
        if stats.get('unlabeled_images'):
            report.append(f"未标注图片：{len(stats.get('unlabeled_images', []))}张（请添加图注）")
        report.append("")
        
        # 审核结论
        report.append("二、审核结论")
        report.append("-"*50)
        report.append(f"审核结果：{'合 格' if review_result.get('是否合格') else '不合格'}")
        report.append(f"最终得分：{review_result.get('总分', 0)}分")
        report.append("")
        
        # 必达项检查
        report.append("三、必达项检查（1/2/4必须达到，3要求框架逻辑正确+图纸必须有）")
        report.append("-"*50)
        must_pass = review_result.get('必达项检查', {})
        if must_pass:
            items = [
                ("题目合格", "1. 题目有具体设计对象"),
                ("字数合格", "2. 字数达到10000字以上"),
                ("框架逻辑合格", "3. 框架逻辑正确（问题→策略→方案）"),
                ("图纸合格", "3. 图纸必须有（设计图+效果图）"),
                ("语言合格", "4. 语言专业通顺，逻辑合理")
            ]
            for key, desc in items:
                result = must_pass.get(key, False)
                report.append(f"   {'✓' if result else '✗'} {desc}")
        report.append("")
        
        # 分项评分
        report.append("四、分项评分与详细评语")
        report.append("-"*50)
        scores = review_result.get('各项评分', {})
        for item, score_info in scores.items():
            report.append(f"\n【{item}】{score_info.get('得分', 0)}分")
            if score_info.get('评语'):
                report.append(f"评语：{score_info['评语']}")
            if score_info.get('问题'):
                for issue in score_info['问题']:
                    report.append(f"问题：{issue}")
        
        report.append("")
        
        # 总体评语
        report.append("五、总体评语")
        report.append("-"*50)
        report.append(review_result.get('总体评语', '无'))
        report.append("")
        
        # 修改建议
        report.append("六、修改建议")
        report.append("-"*50)
        suggestions = review_result.get('修改建议', [])
        for i, suggestion in enumerate(suggestions, 1):
            report.append(f"{i}. {suggestion}")
        
        report.append("")
        report.append("="*70)
        report.append(f"审核时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        report.append("="*70)
        
        return "\n".join(report)
    
    def save_report(self, report: str, output_path: str = None):
        """保存审核报告"""
        if not output_path:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            output_path = f"设计论文审核报告_{timestamp}.txt"
        
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(report)
        
        self.logger.info(f"审核报告已保存至：{output_path}")
        return output_path
    
    def review_thesis(self, file_path: str) -> Dict:
        """审核单篇设计类论文"""
        self.logger.info(f"开始审核设计类论文：{file_path}")
        
        thesis_text, thesis_info = self.extract_text_from_docx(file_path)
        if not thesis_text:
            self.logger.error("论文内容提取失败")
            return None
        
        doc = Document(file_path)
        title = self.extract_title_from_doc(doc)
        
        self.logger.info(f"论文题目：{title}")
        self.logger.info(f"字数：{thesis_info['stats']['word_count']}字")
        self.logger.info(f"设计图：{thesis_info['stats'].get('design_image_count', 0)}张，效果图：{thesis_info['stats'].get('effect_image_count', 0)}张")
        
        self.logger.info("正在使用DeepSeek进行审核...")
        thesis_info['_source_path'] = file_path
        review_result = self.review_with_deepseek(thesis_text, thesis_info, title)
        
        report = self.generate_report(review_result, thesis_info, title)
        self.save_report(report)
        
        self.logger.info(f"审核完成！")
        self.logger.info(f"结果：{'合格' if review_result.get('是否合格') else '不合格'}")
        self.logger.info(f"得分：{review_result.get('总分')}分")
        
        return {
            "result": review_result,
            "report": report,
            "thesis_info": thesis_info
        }


def main():
    """主函数"""
    print("="*70)
    print("无锡太湖学院设计类本科毕业论文审核程序")
    print("（适用于建筑、环境设计、艺术设计等专业）")
    print("="*70)
    print()
    print("【审核说明】")
    print("1. 题目：必须有具体设计对象（建筑/物体/环境氛围）")
    print("2. 字数：10000字以上")
    print("3. 框架逻辑：必须正确（发现问题→提出策略→设计方案）")
    print("   图纸：必须有（设计图+效果图，数量不论，有即可）")
    print("4. 语言：专业通顺，逻辑合理")
    print("5. 图注要求：每张图纸下方请标注图注，如'图4-1 平面图'")
    print("="*70)
    print()
    
    api_key = os.environ.get("DEEPSEEK_API_KEY", "")
    
    while True:
        file_path = input("请输入论文文件路径（支持.docx格式）：").strip()
        if not file_path:
            print("错误：必须指定论文文件路径")
            continue
        if not os.path.exists(file_path):
            print(f"错误：文件不存在 - {file_path}")
            continue
        if not file_path.endswith('.docx'):
            confirm = input("警告：程序主要支持.docx格式，是否继续？(y/n)：").strip().lower()
            if confirm != 'y':
                continue
        break
    
    print("\n正在初始化审核程序...")
    reviewer = ThesisDesignReviewer(api_key=api_key)
    
    print("\n开始审核，请稍候（约60秒）...")
    result = reviewer.review_thesis(file_path)
    
    if result:
        print("\n" + "="*70)
        print("审核完成！")
        print(f"审核结果：{'合 格' if result['result'].get('是否合格') else '不合格'}")
        print(f"最终得分：{result['result'].get('总分')}分")
        print("\n详细报告已保存到文件：设计论文审核报告_时间戳.txt")
        print("="*70)
    else:
        print("\n审核失败，请检查日志。")


if __name__ == "__main__":
    main()